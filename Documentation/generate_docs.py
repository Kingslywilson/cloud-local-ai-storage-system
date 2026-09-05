import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY, TA_RIGHT

def generate_docx(output_path):
    doc = Document()

    # Document Title
    title = doc.add_heading("Gradious AI-Enhanced Cloud File Storage & Sharing Platform", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(30, 58, 138)
        run.font.size = Pt(22)
        run.font.bold = True

    sub = doc.add_paragraph("Comprehensive Project Technical Documentation & Architecture Report")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in sub.runs:
        run.font.color.rgb = RGBColor(71, 85, 105)
        run.font.size = Pt(12)
        run.font.italic = True

    doc.add_paragraph()

    # Section 1: Problem Statement
    h1 = doc.add_heading("1. Problem Statement", level=1)
    p1 = doc.add_paragraph(
        "Traditional cloud file storage platforms often store user assets as opaque binary blobs without understanding "
        "their underlying contents or tracking user security risks. Users struggle to organize, summarize, and extract key insights "
        "from large documents, code repositories, or media files without manual analysis. Furthermore, existing cloud platforms "
        "frequently suffer from insecure file-sharing mechanics (e.g., exposing sequential database IDs) and lack automated, "
        "real-time detection of suspicious upload behaviors or malicious payloads."
    )

    # Section 2: Approach & Solution Strategy
    h2 = doc.add_heading("2. Approach & Solution Strategy", level=1)
    p2 = doc.add_paragraph(
        "To solve these challenges, Gradious introduces an AI-native cloud storage and sharing ecosystem combining three core pillars:"
    )
    bp1 = doc.add_paragraph(style='List Bullet')
    r = bp1.add_run("Full-Stack Cloud Storage Core: ")
    r.bold = True
    bp1.add_run("FastAPI backend with MySQL database providing secure JWT authentication, folder hierarchy, unpredictable UUID share links, storage quota tracking, and comprehensive activity logging.")
    
    bp2 = doc.add_paragraph(style='List Bullet')
    r = bp2.add_run("Machine Learning Module (Random Forest): ")
    r.bold = True
    bp2.add_run("Scikit-learn Random Forest classifiers for automated multi-class file classification (Document, Image, Video, Audio, Code) and anomaly detection for identifying suspicious upload behaviors.")

    bp3 = doc.add_paragraph(style='List Bullet')
    r = bp3.add_run("Generative AI Module (Google Gemini 3.6 Flash): ")
    r.bold = True
    bp3.add_run("Automated document text extraction (PDF, DOCX, TXT, CSV, JSON, MD, Code) and direct multimodal analysis for media assets to produce smart summaries, descriptions, tags, and content insights.")

    # Section 3: System Architecture
    h3 = doc.add_heading("3. System Architecture", level=1)
    doc.add_paragraph(
        "The application utilizes a clean multi-tier architecture separating the client UI, application API server, "
        "relational database, machine learning inference engines, and external generative AI services:"
    )
    
    # Table of Architecture Layers
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Layer"
    hdr_cells[1].text = "Technologies"
    hdr_cells[2].text = "Key Responsibilities"
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True

    layers = [
        ("Frontend Layer", "HTML5, Vanilla CSS3 (Glassmorphism), ES6+ JavaScript", "Responsive multi-page UI, upload forms, AI details display, storage dashboard, sharing modal."),
        ("Backend Layer", "Python 3.11, FastAPI, Uvicorn, Pydantic v2", "REST APIs, JWT token authorization, file I/O sanitization, path traversal prevention, CORS middleware."),
        ("Database Layer", "MySQL 8.0+, SQLAlchemy ORM, PyMySQL", "Relational persistence for users, files, folders, shares, AI analysis metadata, and activity logs."),
        ("Machine Learning Engine", "Scikit-learn, Random Forest, Joblib, Pandas", "Model loading, feature encoding, real-time file category classification & suspicious upload score scoring."),
        ("Generative AI Module", "Google Gemini API (gemini-3.6-flash), PyPDF, Python-Docx", "Document parsing, multimodal prompt execution, structured JSON generation for summary & smart tags.")
    ]

    for layer, tech, resp in layers:
        row_cells = table.add_row().cells
        row_cells[0].text = layer
        row_cells[1].text = tech
        row_cells[2].text = resp

    doc.add_paragraph()

    # Section 4: Tools & Technologies Used
    h4 = doc.add_heading("4. Tools & Technologies Used", level=1)
    tools = [
        ("Programming Languages", "Python 3.11, JavaScript (ES6+), HTML5, CSS3, SQL"),
        ("Backend Framework", "FastAPI, Starlette, Uvicorn ASGI Server"),
        ("Database & ORM", "MySQL 8.0, SQLAlchemy 2.0, PyMySQL"),
        ("Machine Learning", "Scikit-learn (RandomForestClassifier), Joblib, Pandas, NumPy, Matplotlib, Seaborn"),
        ("Generative AI", "Google GenAI SDK (google-genai), Gemini 3.6 Flash LLM"),
        ("Security & Auth", "Passlib (bcrypt), PyJWT / python-jose, Secrets (UUID4)"),
        ("Development Tools", "VS Code, Git, Jupyter Notebook, ReportLab, Python-Docx")
    ]
    for category, detail in tools:
        p = doc.add_paragraph()
        r = p.add_run(f"• {category}: ")
        r.bold = True
        p.add_run(detail)

    # Section 5: Machine Learning Implementation Details
    h5 = doc.add_heading("5. Machine Learning Implementation Details", level=1)
    doc.add_paragraph(
        "The project implements two distinct Random Forest Classifier models saved as serialized joblib binaries:"
    )
    doc.add_paragraph(
        "1. File Category Classifier: Trained on metadata extracted from thousands of files to categorize uploads into "
        "Document, Image, Video, Audio, or Code. Features include file extension, MIME type patterns, and byte size characteristics."
    )
    doc.add_paragraph(
        "2. Suspicious Upload Anomaly Detector: Trained on user behavioral patterns (upload frequency, night-time upload flags, "
        "file size anomalies, unknown extensions, suspicious keyword indicators). Evaluates uploads in real time to assign safety confidence scores."
    )

    # Section 6: Generative AI Implementation Details
    h6 = doc.add_heading("6. Generative AI Implementation Details", level=1)
    doc.add_paragraph(
        "The GenAI module extracts content from uploaded documents (using PyPDF, python-docx, or raw text streams) "
        "or passes image/audio/video files directly to Google Gemini 3.6 Flash. Prompts are structured to return JSON with:"
    )
    doc.add_paragraph("• Smart Overview & Summary (concise summary of document contents)")
    doc.add_paragraph("• Detailed File Description (technical breakdown & purpose)")
    doc.add_paragraph("• Smart Categorical Tags (3-6 descriptive metadata tags)")
    doc.add_paragraph("• Key Content Insights (bullet points detailing critical facts, formulas, or takeaways)")

    # Save docx
    doc.save(output_path)
    print(f"Successfully generated DOCX documentation at: {output_path}")

def generate_pdf(output_path):
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=40, leftMargin=40,
        topMargin=40, bottomMargin=40
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1e3a8a'),
        alignment=TA_CENTER,
        spaceAfter=6
    )

    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#475569'),
        alignment=TA_CENTER,
        spaceAfter=15
    )

    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Heading2'],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#1e3a8a'),
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'Body',
        parent=styles['BodyText'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=8
    )

    bullet_style = ParagraphStyle(
        'Bullet',
        parent=body_style,
        leftIndent=15,
        spaceAfter=4
    )

    story = []

    # Title
    story.append(Paragraph("Gradious AI-Enhanced Cloud File Storage Platform", title_style))
    story.append(Paragraph("Project Technical Documentation & Architecture Report", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceAfter=15))

    # 1. Problem Statement
    story.append(Paragraph("1. Problem Statement", h1_style))
    story.append(Paragraph(
        "Traditional cloud file storage platforms often store user assets as opaque binary blobs without understanding "
        "their underlying contents or tracking user security risks. Users struggle to organize, summarize, and extract key insights "
        "from large documents, code repositories, or media files without manual analysis. Furthermore, existing cloud platforms "
        "frequently suffer from insecure file-sharing mechanics (e.g., exposing sequential database IDs) and lack automated, "
        "real-time detection of suspicious upload behaviors or malicious payloads.",
        body_style
    ))

    # 2. Approach & Solution Strategy
    story.append(Paragraph("2. Approach & Solution Strategy", h1_style))
    story.append(Paragraph(
        "To solve these challenges, Gradious introduces an AI-native cloud storage and sharing ecosystem combining three core pillars:",
        body_style
    ))
    story.append(Paragraph("• <b>Full-Stack Storage Core:</b> FastAPI backend with MySQL database providing secure JWT authentication, folder hierarchy, unpredictable UUID share links, storage quota tracking, and comprehensive activity logging.", bullet_style))
    story.append(Paragraph("• <b>Machine Learning Module (Random Forest):</b> Scikit-learn Random Forest classifiers for automated multi-class file classification (Document, Image, Video, Audio, Code) and anomaly detection for identifying suspicious upload behaviors.", bullet_style))
    story.append(Paragraph("• <b>Generative AI Module (Google Gemini 3.6 Flash):</b> Automated document text extraction (PDF, DOCX, TXT, CSV, JSON, MD, Code) and direct multimodal analysis for media assets to produce smart summaries, descriptions, tags, and content insights.", bullet_style))

    # 3. System Architecture Table
    story.append(Paragraph("3. System Architecture", h1_style))
    
    table_data = [
        [Paragraph("<b>Layer</b>", body_style), Paragraph("<b>Technologies</b>", body_style), Paragraph("<b>Key Responsibilities</b>", body_style)],
        [Paragraph("Frontend", body_style), Paragraph("HTML5, Vanilla CSS3, JavaScript ES6+", body_style), Paragraph("Responsive UI, upload forms, AI details display, storage dashboard, link sharing.", body_style)],
        [Paragraph("Backend", body_style), Paragraph("Python 3.11, FastAPI, Uvicorn", body_style), Paragraph("REST APIs, JWT authorization, file I/O sanitization, path traversal prevention.", body_style)],
        [Paragraph("Database", body_style), Paragraph("MySQL 8.0, SQLAlchemy, PyMySQL", body_style), Paragraph("Relational persistence for users, files, folders, shares, AI metadata, activity logs.", body_style)],
        [Paragraph("ML Engine", body_style), Paragraph("Scikit-learn, Random Forest, Joblib", body_style), Paragraph("Feature encoding, real-time file classification & suspicious upload scoring.", body_style)],
        [Paragraph("GenAI Module", body_style), Paragraph("Google Gemini 3.6 Flash, PyPDF, Docx", body_style), Paragraph("Document parsing, multimodal LLM execution, structured JSON summaries & tags.", body_style)]
    ]

    t = Table(table_data, colWidths=[80, 150, 280])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t)
    story.append(Spacer(1, 10))

    # 4. Tools & Technologies
    story.append(Paragraph("4. Tools & Technologies Used", h1_style))
    tools = [
        ("Programming Languages", "Python 3.11, JavaScript (ES6+), HTML5, CSS3, SQL"),
        ("Backend Framework", "FastAPI, Starlette, Uvicorn ASGI Server"),
        ("Database & ORM", "MySQL 8.0, SQLAlchemy 2.0, PyMySQL"),
        ("Machine Learning", "Scikit-learn (RandomForestClassifier), Joblib, Pandas, NumPy, Matplotlib, Seaborn"),
        ("Generative AI", "Google GenAI SDK (google-genai), Gemini 3.6 Flash LLM"),
        ("Security & Auth", "Passlib (bcrypt), PyJWT / python-jose, Secrets (UUID4)"),
        ("Development Tools", "VS Code, Git, Jupyter Notebook, ReportLab, Python-Docx")
    ]
    for category, detail in tools:
        story.append(Paragraph(f"• <b>{category}:</b> {detail}", bullet_style))

    # 5. Machine Learning Implementation Details
    story.append(Paragraph("5. Machine Learning Implementation Details", h1_style))
    story.append(Paragraph(
        "The project implements two distinct Random Forest Classifier models saved as serialized joblib binaries:<br/>"
        "1. <b>File Category Classifier:</b> Trained on metadata extracted from thousands of files to categorize uploads into Document, Image, Video, Audio, or Code.<br/>"
        "2. <b>Suspicious Upload Anomaly Detector:</b> Trained on user behavioral patterns (upload frequency, night-time upload flags, file size anomalies, unknown extensions, suspicious keyword indicators). Evaluates uploads in real time to assign safety confidence scores.",
        body_style
    ))

    # 6. Generative AI Implementation Details
    story.append(Paragraph("6. Generative AI Implementation Details", h1_style))
    story.append(Paragraph(
        "The GenAI module extracts content from uploaded documents (using PyPDF, python-docx, or raw text streams) "
        "or passes image/audio/video files directly to Google Gemini 3.6 Flash. Prompts are structured to return JSON with "
        "Smart Overview, Detailed File Description, Smart Categorical Tags, and Key Content Insights.",
        body_style
    ))

    doc.build(story)
    print(f"Successfully generated PDF documentation at: {output_path}")

if __name__ == "__main__":
    docs_dir = Path("Documentation")
    docs_dir.mkdir(exist_ok=True)
    generate_docx(docs_dir / "Project_Documentation.docx")
    generate_pdf(docs_dir / "Project_Documentation.pdf")
